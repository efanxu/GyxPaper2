from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_ROOT = PROJECT_ROOT / "custom_models" / "src"
sys.path.insert(0, str(SRC_ROOT))

from benchmark_v2.artifacts import atomic_write_json, validate_run  # noqa: E402
from benchmark_v2.experiments.e5_common_loss.active_scope import (  # noqa: E402
    ActiveScopeError,
    active_scope_path,
    load_active_scope_pointer,
)
from benchmark_v2.experiments.e5_common_loss.a8_reference import (  # noqa: E402
    validate_a8_reference,
)
from benchmark_v2.experiments.e5_common_loss.loss_profile import (  # noqa: E402
    CLI_PROFILE_ID,
    get_profile_metadata,
)
from benchmark_v2.experiments.e5_common_loss.runner import (  # noqa: E402
    canonical_base_model_config_hash,
)
from benchmark_v2.experiments.e5_common_loss.scope27_contract import (  # noqa: E402
    E5_SCOPE27_EVALUATE_ONLY_MODELS,
    E5_SCOPE27_EXCLUDED_MODELS,
    E5_SCOPE27_ID,
    E5_SCOPE27_REFERENCE_MODELS,
    E5_SCOPE27_TRAINABLE_MODELS,
    TRAINING_PROFILE_ID,
)
from benchmark_v2.protocol import load_protocol  # noqa: E402
from benchmark_v2.registry import load_registry  # noqa: E402
from benchmark_v2.model_source_identity import (  # noqa: E402
    MODEL_SOURCE_IDENTITY_SCHEMA_VERSION,
    ModelSourceIdentityError,
    canonical_model_source_identity,
)
from benchmark_v2.precision import expected_model_precision_identity  # noqa: E402
from benchmark_v2.training_profiles import load_training_profile  # noqa: E402


DEFAULT_MANIFEST = (
    PROJECT_ROOT
    / "custom_models"
    / "docs"
    / "benchmark_v2"
    / "E5"
    / "E5_SCOPE27_VARIANT_MANIFEST.json"
)
EXPECTED_OUTPUT_ROOT = (
    PROJECT_ROOT
    / "custom_models"
    / "results"
    / "benchmark_v2"
    / "common_loss_architecture_seed2026"
)
REQUIRED_METRICS = ("Score", "MAE", "RMSE", "R2")
HORIZONS = (3, 6, 10)
SOURCE_IDENTITY_SCHEMA_VERSION = MODEL_SOURCE_IDENTITY_SCHEMA_VERSION


class ScopeGateError(RuntimeError):
    pass


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def combined_hash(paths: Iterable[Path]) -> str:
    digest = hashlib.sha256()
    for path in paths:
        digest.update(path.name.encode("utf-8"))
        digest.update(bytes.fromhex(sha256_file(path)))
    return digest.hexdigest()


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def load_manifest(path: str | Path) -> dict[str, Any]:
    pointer = load_active_scope_pointer()
    manifest_path = Path(path).resolve()
    expected_manifest = active_scope_path(pointer, "manifest").resolve()
    if manifest_path != expected_manifest:
        raise ScopeGateError(
            "The active scope pointer does not authorize this manifest; "
            "legacy/superseded manifests are not current scope27 inputs."
        )
    payload = load_json(manifest_path)
    if not isinstance(payload, dict):
        raise ScopeGateError("Manifest root must be a JSON object.")
    validate_manifest(payload)
    return payload


def _entry_sets(manifest: Mapping[str, Any]) -> dict[str, tuple[str, ...]]:
    entries = list(manifest["entries"])
    return {
        "trainable": tuple(
            row["model_id"]
            for row in entries
            if row["entry_type"] == "TRAIN_COMMON_LOSS"
        ),
        "evaluate_only": tuple(
            row["model_id"]
            for row in entries
            if row["entry_type"] == "EVALUATE_ONLY_COMMON_LOSS_DIAGNOSTIC"
        ),
        "reference": tuple(
            row["model_id"]
            for row in entries
            if row["entry_type"] == "REFERENCE_ONLY_FORMAL_A8"
        ),
    }


def validate_manifest(manifest: Mapping[str, Any]) -> None:
    if manifest.get("scope_id") != E5_SCOPE27_ID:
        raise ScopeGateError("Active E5 scope identity mismatch.")
    if manifest.get("training_profile_id") != TRAINING_PROFILE_ID:
        raise ScopeGateError("Training profile identity mismatch.")
    if manifest.get("cli_profile_id") != CLI_PROFILE_ID:
        raise ScopeGateError("E5 CLI profile identity mismatch.")
    if manifest.get("output_root") != (
        "custom_models/results/benchmark_v2/common_loss_architecture_seed2026"
    ):
        raise ScopeGateError("Formal output root mismatch.")
    if manifest.get("source_identity_schema_version") != SOURCE_IDENTITY_SCHEMA_VERSION:
        raise ScopeGateError("Model source identity schema mismatch.")
    entries = list(manifest.get("entries", []))
    counts = manifest.get("counts", {})
    expected_counts = {
        "trainable": 24,
        "evaluate_only": 2,
        "a8_reference": 1,
        "total_evidence": 27,
    }
    if counts != expected_counts or len(entries) != 27:
        raise ScopeGateError(f"E5 scope27 count mismatch: {counts}")
    entry_sets = _entry_sets(manifest)
    expected_sets = {
        "trainable": E5_SCOPE27_TRAINABLE_MODELS,
        "evaluate_only": E5_SCOPE27_EVALUATE_ONLY_MODELS,
        "reference": E5_SCOPE27_REFERENCE_MODELS,
    }
    if entry_sets != expected_sets:
        raise ScopeGateError(
            f"E5 scope27 entry order/content mismatch: {entry_sets}"
        )
    model_ids = [row["model_id"] for row in entries]
    run_ids = [row["e5_run_id"] for row in entries]
    if len(set(model_ids)) != 27 or len(set(run_ids)) != 27:
        raise ScopeGateError("Model ids and run ids must be unique.")
    if set(model_ids) & set(E5_SCOPE27_EXCLUDED_MODELS):
        raise ScopeGateError("Excluded model is present in the active manifest.")
    if [row.get("ordinal") for row in entries] != list(range(1, 28)):
        raise ScopeGateError("Manifest ordinals must be exactly 1..27.")

    protocol = load_protocol()
    if manifest.get("benchmark_protocol_hash") != protocol.protocol_hash:
        raise ScopeGateError("Benchmark Protocol identity mismatch.")
    profile = load_training_profile(TRAINING_PROFILE_ID)
    if profile is None:
        raise ScopeGateError("Frozen batch4 training profile is unavailable.")
    identity = profile.identity()
    for key, expected in (
        ("training_profile_hash", profile.profile_hash),
        ("train_batch_size", 4),
        ("val_batch_size", 4),
        ("test_batch_size", 4),
    ):
        actual = (
            manifest.get("batch_identity", {}).get(key)
            if key.endswith("_batch_size")
            else manifest.get(key)
        )
        if actual != expected:
            raise ScopeGateError(
                f"Frozen batch4 identity mismatch for {key}: {actual!r}"
            )
    if any(identity[key] != 4 for key in (
        "train_batch_size",
        "val_batch_size",
        "test_batch_size",
    )):
        raise ScopeGateError("The active training profile is not batch=4.")

    loss = get_profile_metadata(CLI_PROFILE_ID)
    expected_loss = manifest.get("loss_identity", {})
    for key in (
        "loss_id",
        "loss_source_hash",
        "loss_profile_hash",
        "e5_common_loss_protocol_hash",
    ):
        if expected_loss.get(key) != loss[key]:
            raise ScopeGateError(f"E5 loss identity mismatch for {key}.")

    registry = load_registry()
    for entry in entries:
        if entry["model_id"] in E5_SCOPE27_REFERENCE_MODELS:
            continue
        try:
            source_identity = canonical_model_source_identity(entry["model_id"])
        except (ModelSourceIdentityError, KeyError, OSError) as exc:
            raise ScopeGateError(
                f"Model source closure cannot be validated: {entry['model_id']}: {exc}"
            ) from exc
        if (
            entry.get("base_model_source_hash")
            != source_identity["canonical_combined_hash"]
            or entry.get("base_model_source_closure_hash")
            != source_identity["canonical_combined_hash"]
        ):
            raise ScopeGateError(
                f"Model source identity mismatch: {entry['model_id']}"
            )
        stored_identity = entry.get("source_identity")
        if not isinstance(stored_identity, Mapping):
            raise ScopeGateError(
                f"Model source closure manifest is missing: {entry['model_id']}"
            )
        if stored_identity.get("source_identity_schema_version") != SOURCE_IDENTITY_SCHEMA_VERSION:
            raise ScopeGateError(
                f"Model source closure schema mismatch: {entry['model_id']}"
            )
        if stored_identity.get("canonical_combined_hash") != source_identity[
            "canonical_combined_hash"
        ] or stored_identity.get("source_closure_files") != source_identity[
            "source_closure_files"
        ]:
            raise ScopeGateError(
                f"Model source closure records mismatch: {entry['model_id']}"
            )
        if (
            canonical_base_model_config_hash(entry["model_id"])
            != entry["base_model_config_hash"]
        ):
            raise ScopeGateError(
                f"Model config identity mismatch: {entry['model_id']}"
            )
        if entry.get("precision_identity") != expected_model_precision_identity(
            entry["model_id"], TRAINING_PROFILE_ID
        ):
            raise ScopeGateError(
                f"Precision identity mismatch: {entry['model_id']}"
            )


def _entry(manifest: Mapping[str, Any], model_id: str) -> Mapping[str, Any]:
    for row in manifest["entries"]:
        if row["model_id"] == model_id:
            return row
    raise ScopeGateError(f"Model is not in current E5 scope: {model_id}")


def _metric_payloads(paths: list[Path]) -> tuple[dict[int, dict[str, Any]], list[str]]:
    reasons: list[str] = []
    payloads: dict[int, dict[str, Any]] = {}
    for horizon, path in zip(HORIZONS, paths):
        if not path.is_file():
            reasons.append(f"MISSING:{path.name}")
            continue
        payload = load_json(path)
        if payload.get("horizon") != horizon:
            reasons.append(f"HORIZON_MISMATCH:{path.name}")
        for key in REQUIRED_METRICS:
            value = payload.get(key)
            if (
                isinstance(value, bool)
                or not isinstance(value, (int, float))
                or not math.isfinite(float(value))
            ):
                reasons.append(f"NON_FINITE_OR_MISSING:{path.name}:{key}")
        payloads[horizon] = payload
    return payloads, reasons


def _base_run_result(
    entry: Mapping[str, Any], output_root: Path
) -> tuple[dict[str, Any], Path]:
    run_dir = output_root / entry["e5_run_id"]
    return (
        {
            "entry_id": entry["entry_id"],
            "model_id": entry["model_id"],
            "entry_type": entry["entry_type"],
            "run_id": entry["e5_run_id"],
            "run_dir": str(run_dir),
            "found": run_dir.is_dir(),
            "ready": False,
            "reasons": [],
        },
        run_dir,
    )


def _load_common_run(
    manifest: Mapping[str, Any],
    entry: Mapping[str, Any],
    output_root: Path,
    *,
    expected_artifact_profile: str,
    expected_formal_training: bool,
) -> tuple[dict[str, Any], Path, dict[str, Any] | None]:
    result, run_dir = _base_run_result(entry, output_root)
    if not run_dir.is_dir():
        result["reasons"].append("RUN_NOT_FOUND")
        return result, run_dir, None
    try:
        status = load_json(run_dir / "run_status.json")
        effective = load_json(run_dir / "effective_config.json")
        data_signature = load_json(run_dir / "data_signature.json")
        protocol_check = load_json(run_dir / "protocol_check.json")
        artifact_manifest = load_json(run_dir / "artifact_manifest.json")
        validate_run(
            run_dir,
            expected_protocol_hash=manifest["benchmark_protocol_hash"],
            expected_training_batch_profile_id=TRAINING_PROFILE_ID,
            expected_training_batch_profile_hash=manifest[
                "training_profile_hash"
            ],
        )
    except Exception as exc:
        result["reasons"].append(
            f"ARTIFACT_INTEGRITY:{type(exc).__name__}:{exc}"
        )
        return result, run_dir, None

    provenance = effective.get("provenance", {})
    stored_identity = entry.get("source_identity", {})
    expected_precision = entry.get("precision_identity", {})
    checks = (
        (status.get("status") == "COMPLETED", "NOT_COMPLETED"),
        (status.get("run_mode") == "formal", "NOT_FORMAL"),
        (
            status.get("artifact_profile") == expected_artifact_profile,
            "ARTIFACT_PROFILE_MISMATCH",
        ),
        (
            status.get("formal_training") is expected_formal_training,
            "FORMAL_TRAINING_FLAG_MISMATCH",
        ),
        (effective.get("model_id") == entry["model_id"], "MODEL_ID_MISMATCH"),
        (
            effective.get("artifact_profile") == expected_artifact_profile,
            "EFFECTIVE_ARTIFACT_PROFILE_MISMATCH",
        ),
        (
            effective.get("formal_training") is expected_formal_training,
            "EFFECTIVE_FORMAL_TRAINING_FLAG_MISMATCH",
        ),
        (
            effective.get("experiment_profile_id")
            == manifest["experiment_profile_id"],
            "E5_PROFILE_MISMATCH",
        ),
        (
            provenance.get("active_scope_id") == E5_SCOPE27_ID,
            "ACTIVE_SCOPE_ID_MISMATCH",
        ),
        (
            provenance.get("base_model_config_hash")
            == entry["base_model_config_hash"],
            "BASE_MODEL_CONFIG_HASH_MISMATCH",
        ),
        (
            provenance.get("base_model_source_closure_hash")
            == entry["base_model_source_hash"],
            "BASE_MODEL_SOURCE_HASH_MISMATCH",
        ),
        (
            provenance.get("base_model_source_identity_schema_version")
            == SOURCE_IDENTITY_SCHEMA_VERSION,
            "BASE_MODEL_SOURCE_SCHEMA_MISMATCH",
        ),
        (
            provenance.get("base_model_source_closure_files")
            == stored_identity.get("source_closure_files"),
            "BASE_MODEL_SOURCE_CLOSURE_MISMATCH",
        ),
        (
            effective.get("loss", {}).get("id")
            == manifest["loss_identity"]["loss_id"],
            "LOSS_ID_MISMATCH",
        ),
        (
            effective.get("loss", {}).get("source_hash")
            == manifest["loss_identity"]["loss_source_hash"],
            "LOSS_SOURCE_HASH_MISMATCH",
        ),
        (
            effective.get("loss", {}).get("profile_hash")
            == manifest["loss_identity"]["loss_profile_hash"],
            "LOSS_PROFILE_HASH_MISMATCH",
        ),
        (
            provenance.get("e5_common_loss_protocol_hash")
            == manifest["loss_identity"]["e5_common_loss_protocol_hash"],
            "E5_PROTOCOL_HASH_MISMATCH",
        ),
        (
            protocol_check.get("protocol_hash")
            == manifest["benchmark_protocol_hash"],
            "BENCHMARK_PROTOCOL_HASH_MISMATCH",
        ),
        (artifact_manifest.get("model_id") == entry["model_id"], "ARTIFACT_MODEL_ID_MISMATCH"),
        (
            artifact_manifest.get("artifact_profile") == expected_artifact_profile,
            "ARTIFACT_MANIFEST_PROFILE_MISMATCH",
        ),
        (data_signature.get("dataset_id") == "SDWPF", "DATASET_ID_MISMATCH"),
        (data_signature.get("node_count") == 134, "DATASET_NODE_COUNT_MISMATCH"),
        (
            data_signature.get("feature_order_hash")
            == manifest["dataset_identity"]["feature_order_hash"],
            "DATASET_FEATURE_ORDER_MISMATCH",
        ),
        (
            all(
                effective.get(key) == value
                for key, value in expected_precision.items()
            ),
            "PRECISION_IDENTITY_MISMATCH",
        ),
    )
    result["reasons"].extend(reason for passed, reason in checks if not passed)
    return result, run_dir, {
        "status": status,
        "effective": effective,
        "data_signature": data_signature,
        "protocol_check": protocol_check,
        "artifact_manifest": artifact_manifest,
    }


def inspect_trainable_run(
    manifest: Mapping[str, Any],
    entry: Mapping[str, Any],
    output_root: Path,
) -> dict[str, Any]:
    result, run_dir, common = _load_common_run(
        manifest,
        entry,
        output_root,
        expected_artifact_profile="TRAIN",
        expected_formal_training=True,
    )
    if common is None:
        return result
    metrics_paths = [
        run_dir / f"metrics_eval_h{horizon}.json" for horizon in HORIZONS
    ]
    _, metric_reasons = _metric_payloads(metrics_paths)
    result["reasons"].extend(metric_reasons)
    checkpoint = run_dir / "best_checkpoint.pt"
    if not checkpoint.is_file():
        result["reasons"].append("BEST_CHECKPOINT_MISSING")
    if not result["reasons"]:
        result.update(
            {
                "ready": True,
                "checkpoint_sha256": sha256_file(checkpoint),
                "metrics_sha256": combined_hash(metrics_paths),
                "run_status_sha256": sha256_file(run_dir / "run_status.json"),
                "effective_config_sha256": sha256_file(
                    run_dir / "effective_config.json"
                ),
                "artifact_manifest_sha256": sha256_file(
                    run_dir / "artifact_manifest.json"
                ),
            }
        )
    return result


def inspect_evaluate_only_run(
    manifest: Mapping[str, Any],
    entry: Mapping[str, Any],
    output_root: Path,
) -> dict[str, Any]:
    result, run_dir, common = _load_common_run(
        manifest,
        entry,
        output_root,
        expected_artifact_profile="NON_TRAINABLE",
        expected_formal_training=False,
    )
    if common is None:
        return result
    baseline_path = run_dir / "baseline_state.json"
    diagnostic_path = run_dir / "common_loss_diagnostic.json"
    prediction_path = run_dir / "prediction_metadata.json"
    try:
        baseline = load_json(baseline_path)
        diagnostic = load_json(diagnostic_path)
        prediction = load_json(prediction_path)
    except Exception as exc:
        result["reasons"].append(
            f"EVALUATE_ONLY_EVIDENCE:{type(exc).__name__}:{exc}"
        )
        return result
    baseline_checks = (
        (
            baseline.get("schema_version") == "e5_non_trainable_baseline_v1",
            "BASELINE_SCHEMA_MISMATCH",
        ),
        (baseline.get("model_id") == entry["model_id"], "BASELINE_MODEL_ID_MISMATCH"),
        (baseline.get("training_mode") == "EVALUATE_ONLY", "BASELINE_MODE_MISMATCH"),
        (baseline.get("training_loss") == "NOT_APPLICABLE", "BASELINE_TRAINING_LOSS_MISMATCH"),
        (baseline.get("trained_with_common_loss") is False, "BASELINE_TRAINED_FLAG_MISMATCH"),
        (baseline.get("common_loss_evaluation_applied") is True, "BASELINE_DIAGNOSTIC_FLAG_MISMATCH"),
        (baseline.get("best_checkpoint") is None, "BASELINE_CHECKPOINT_NOT_NULL"),
        (baseline.get("best_epoch") is None, "BASELINE_BEST_EPOCH_NOT_NULL"),
    )
    diagnostic_checks = (
        (
            diagnostic.get("schema_version") == "e5_common_loss_diagnostic_v1",
            "DIAGNOSTIC_SCHEMA_MISMATCH",
        ),
        (
            diagnostic.get("model_id") == entry["model_id"],
            "DIAGNOSTIC_MODEL_ID_MISMATCH",
        ),
        (diagnostic.get("loss_id") == "masked_score_aligned_hybrid", "DIAGNOSTIC_LOSS_ID_MISMATCH"),
        (
            diagnostic.get("loss_space") == "normalized Patv_raw",
            "DIAGNOSTIC_LOSS_SPACE_MISMATCH",
        ),
        (diagnostic.get("training_loss") == "NOT_APPLICABLE", "DIAGNOSTIC_TRAINING_LOSS_MISMATCH"),
        (diagnostic.get("trained_with_common_loss") is False, "DIAGNOSTIC_TRAINED_FLAG_MISMATCH"),
        (diagnostic.get("common_loss_evaluation_applied") is True, "DIAGNOSTIC_APPLIED_FLAG_MISMATCH"),
        (
            isinstance(diagnostic.get("value"), (int, float))
            and not isinstance(diagnostic.get("value"), bool)
            and math.isfinite(float(diagnostic["value"])),
            "DIAGNOSTIC_VALUE_NON_FINITE",
        ),
    )
    result["reasons"].extend(
        reason
        for passed, reason in (*baseline_checks, *diagnostic_checks)
        if not passed
    )
    if prediction.get("source_checkpoint") is not None:
        result["reasons"].append("EVALUATE_ONLY_SOURCE_CHECKPOINT_NOT_NULL")
    metrics_paths = [
        run_dir / f"metrics_eval_h{horizon}.json" for horizon in HORIZONS
    ]
    _, metric_reasons = _metric_payloads(metrics_paths)
    result["reasons"].extend(metric_reasons)
    if not result["reasons"]:
        result.update(
            {
                "ready": True,
                "checkpoint_sha256": None,
                "baseline_state_sha256": sha256_file(baseline_path),
                "common_loss_diagnostic_sha256": sha256_file(diagnostic_path),
                "metrics_sha256": combined_hash(metrics_paths),
                "run_status_sha256": sha256_file(run_dir / "run_status.json"),
                "effective_config_sha256": sha256_file(
                    run_dir / "effective_config.json"
                ),
                "artifact_manifest_sha256": sha256_file(
                    run_dir / "artifact_manifest.json"
                ),
            }
        )
    return result


def inspect_run(
    manifest: Mapping[str, Any],
    entry: Mapping[str, Any],
    output_root: Path,
) -> dict[str, Any]:
    if entry["entry_type"] == "TRAIN_COMMON_LOSS":
        return inspect_trainable_run(manifest, entry, output_root)
    if entry["entry_type"] == "EVALUATE_ONLY_COMMON_LOSS_DIAGNOSTIC":
        return inspect_evaluate_only_run(manifest, entry, output_root)
    raise ScopeGateError(f"Unsupported run entry type: {entry['entry_type']}")


def inspect_a8(
    manifest: Mapping[str, Any], entry: Mapping[str, Any]
) -> dict[str, Any]:
    expected = manifest["a8_reference_identity"]
    result = validate_a8_reference()
    reasons: list[str] = []
    checks = (
        (result.get("status") == "VALID", "A8_REFERENCE_INVALID"),
        (result.get("reference_id") == entry["e5_run_id"], "A8_ID_MISMATCH"),
        (
            result.get("source_checkpoint_sha256")
            == expected["checkpoint_sha256"],
            "A8_CHECKPOINT_HASH_MISMATCH",
        ),
        (
            result.get("source_metrics_sha256") == expected["metrics_sha256"],
            "A8_METRICS_HASH_MISMATCH",
        ),
        (
            result.get("source_config_sha256") == expected["config_sha256"],
            "A8_CONFIG_HASH_MISMATCH",
        ),
        (
            result.get("source_protocol_evidence_sha256")
            == expected["protocol_evidence_sha256"],
            "A8_PROTOCOL_EVIDENCE_HASH_MISMATCH",
        ),
        (
            result.get("source_protocol_hash")
            == manifest["benchmark_protocol_hash"],
            "A8_PROTOCOL_HASH_MISMATCH",
        ),
        (
            result.get("source_loss_id")
            == manifest["loss_identity"]["loss_id"],
            "A8_LOSS_ID_MISMATCH",
        ),
        (result.get("retrained_in_e5") is False, "A8_MUST_BE_READ_ONLY"),
    )
    reasons.extend(reason for passed, reason in checks if not passed)
    return {
        "entry_id": entry["entry_id"],
        "model_id": entry["model_id"],
        "entry_type": entry["entry_type"],
        "run_id": entry["e5_run_id"],
        "run_dir": result.get("source_absolute_or_resolved_path"),
        "found": not bool(result.get("missing")),
        "ready": not reasons,
        "reasons": reasons,
        "checkpoint_sha256": result.get("source_checkpoint_sha256"),
        "metrics_sha256": result.get("source_metrics_sha256"),
        "effective_config_sha256": result.get("source_config_sha256"),
        "protocol_evidence_sha256": result.get(
            "source_protocol_evidence_sha256"
        ),
        "reference": result,
    }


def build_readiness(
    manifest: Mapping[str, Any], output_root: Path
) -> dict[str, Any]:
    expected = EXPECTED_OUTPUT_ROOT.resolve()
    actual = output_root.resolve()
    if actual != expected:
        raise ScopeGateError(
            f"Output root mismatch: actual={actual}, expected={expected}"
        )
    rows = []
    for entry in manifest["entries"]:
        if entry["entry_type"] == "REFERENCE_ONLY_FORMAL_A8":
            rows.append(inspect_a8(manifest, entry))
        else:
            rows.append(inspect_run(manifest, entry, actual))
    by_type = {
        "trainable": [
            row for row in rows if row["entry_type"] == "TRAIN_COMMON_LOSS"
        ],
        "evaluate_only": [
            row
            for row in rows
            if row["entry_type"]
            == "EVALUATE_ONLY_COMMON_LOSS_DIAGNOSTIC"
        ],
        "a8_reference": [
            row
            for row in rows
            if row["entry_type"] == "REFERENCE_ONLY_FORMAL_A8"
        ],
    }
    counts = {
        key: {
            "ready": sum(row["ready"] for row in selected),
            "expected": len(selected),
        }
        for key, selected in by_type.items()
    }
    counts["total_evidence"] = {
        "ready": sum(row["ready"] for row in rows),
        "expected": 27,
    }
    return {
        "schema_version": "e5_scope27_readiness_v1",
        "scope_id": E5_SCOPE27_ID,
        "status": (
            "READY"
            if counts["total_evidence"]["ready"] == 27
            else "NOT_READY"
        ),
        "require_complete": True,
        "counts": counts,
        "output_root": str(actual),
        "generated_at": utc_now(),
        "entries": rows,
    }


def write_readiness_outputs(
    readiness: Mapping[str, Any],
    report_path: Path,
    evidence_path: Path,
) -> None:
    atomic_write_json(report_path, readiness)
    evidence = {
        "schema_version": "e5_scope27_evidence_manifest_v1",
        "scope_id": E5_SCOPE27_ID,
        "status": readiness["status"],
        "require_complete": True,
        "generated_at": readiness["generated_at"],
        "entries": [
            {
                key: row.get(key)
                for key in (
                    "entry_id",
                    "model_id",
                    "entry_type",
                    "run_id",
                    "run_dir",
                    "ready",
                    "checkpoint_sha256",
                    "baseline_state_sha256",
                    "common_loss_diagnostic_sha256",
                    "metrics_sha256",
                    "run_status_sha256",
                    "effective_config_sha256",
                    "artifact_manifest_sha256",
                    "protocol_evidence_sha256",
                )
            }
            for row in readiness["entries"]
        ],
    }
    atomic_write_json(evidence_path, evidence)


def _aggregate_rows(
    manifest: Mapping[str, Any], readiness: Mapping[str, Any]
) -> list[dict[str, Any]]:
    entries = {row["model_id"]: row for row in manifest["entries"]}
    result: list[dict[str, Any]] = []
    for evidence in readiness["entries"]:
        entry = entries[evidence["model_id"]]
        if entry["entry_type"] == "REFERENCE_ONLY_FORMAL_A8":
            reference = evidence["reference"]
            metrics_paths = [
                Path(path) for path in reference["source_metrics_paths"]
            ]
            effective = load_json(Path(reference["source_config_path"]))
            training_status = "REFERENCE_ONLY"
        else:
            run_dir = Path(evidence["run_dir"])
            metrics_paths = [
                run_dir / f"metrics_eval_h{horizon}.json"
                for horizon in HORIZONS
            ]
            effective = load_json(run_dir / "effective_config.json")
            training_status = load_json(run_dir / "run_status.json")["status"]
        metrics, reasons = _metric_payloads(metrics_paths)
        if reasons:
            raise ScopeGateError(
                f"Refusing non-finite/incomplete aggregate for "
                f"{entry['model_id']}: {reasons}"
            )
        row: dict[str, Any] = {
            "model": entry["display_name"],
            "model_id": entry["model_id"],
            "entry_type": entry["entry_type"],
            "training_mode": (
                "REFERENCE_ONLY"
                if entry["entry_type"] == "REFERENCE_ONLY_FORMAL_A8"
                else ("EVALUATE_ONLY" if entry["command"] == "evaluate-only" else "TRAIN")
            ),
            "loss_id": entry["loss_id"],
            "training_status": training_status,
            "checkpoint": (
                "best_checkpoint.pt"
                if entry["entry_type"] == "TRAIN_COMMON_LOSS"
                else None
            ),
            "best_epoch": effective.get("best_epoch"),
            "parameter_count": effective.get("parameter_count"),
            "trainable_parameter_count": effective.get(
                "trainable_parameter_count"
            ),
            "active_scope_id": E5_SCOPE27_ID,
            "training_batch_profile_id": (
                None
                if entry["entry_type"] == "REFERENCE_ONLY_FORMAL_A8"
                else TRAINING_PROFILE_ID
            ),
        }
        for horizon in HORIZONS:
            for metric in REQUIRED_METRICS:
                row[f"H{horizon}_{metric}"] = metrics[horizon][metric]
        row["average_Score"] = sum(
            float(metrics[horizon]["Score"]) for horizon in HORIZONS
        ) / len(HORIZONS)
        result.append(row)
    return result


def aggregate(
    manifest: Mapping[str, Any],
    output_root: Path,
    *,
    require_complete: bool,
    report_path: Path,
    evidence_path: Path,
) -> dict[str, Any]:
    if not require_complete:
        raise ScopeGateError("Final E5 aggregation requires --require-complete.")
    readiness = build_readiness(manifest, output_root)
    write_readiness_outputs(readiness, report_path, evidence_path)
    if readiness["status"] != "READY":
        ready = readiness["counts"]["total_evidence"]["ready"]
        raise ScopeGateError(f"E5_SCOPE27_NOT_READY:{ready}/27")
    rows = _aggregate_rows(manifest, readiness)
    if len(rows) != 27:
        raise ScopeGateError("Aggregate row count must be exactly 27.")

    fieldnames = list(rows[0])
    csv_path = output_root / "common_loss_architecture_scope27_seed2026.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    md_path = output_root / "COMMON_LOSS_ARCHITECTURE_SCOPE27_SEED2026.md"
    lines = [
        "# E5 common-loss architecture: batch4 scope27 seed2026",
        "",
        "| model | mode | checkpoint | H3 Score | H6 Score | H10 Score | average Score |",
        "|---|---|---|---:|---:|---:|---:|",
    ]
    for row in rows:
        lines.append(
            f"| {row['model']} | {row['training_mode']} | "
            f"{row['checkpoint'] or 'null'} | "
            f"{row['H3_Score']:.6f} | {row['H6_Score']:.6f} | "
            f"{row['H10_Score']:.6f} | {row['average_Score']:.6f} |"
        )
    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    audit_path = output_root / "common_loss_architecture_scope27_audit.json"
    atomic_write_json(
        audit_path,
        {
            "status": "PASS",
            "scope_id": E5_SCOPE27_ID,
            "require_complete": True,
            "row_count": 27,
            "readiness": readiness["counts"],
            "generated_at": utc_now(),
        },
    )
    xlsx_path = output_root / "COMMON_LOSS_ARCHITECTURE_SCOPE27_SEED2026.xlsx"
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise ScopeGateError(
            "openpyxl is required for the complete E5 aggregate."
        ) from exc
    workbook = Workbook()
    trained = workbook.active
    trained.title = "Trainable structures"
    references = workbook.create_sheet("Evaluate-only and A8")
    all_rows = workbook.create_sheet("All 27")
    selections = (
        (
            trained,
            [row for row in rows if row["entry_type"] == "TRAIN_COMMON_LOSS"],
        ),
        (
            references,
            [row for row in rows if row["entry_type"] != "TRAIN_COMMON_LOSS"],
        ),
        (all_rows, rows),
    )
    for sheet, selected in selections:
        sheet.append(fieldnames)
        for row in selected:
            sheet.append([row.get(field) for field in fieldnames])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    workbook.save(xlsx_path)
    return {
        "status": "PASS",
        "scope_id": E5_SCOPE27_ID,
        "row_count": 27,
        "files": [
            str(xlsx_path),
            str(md_path),
            str(csv_path),
            str(audit_path),
        ],
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="e5_batch4_scope27_gate")
    parser.add_argument("--manifest", default=str(DEFAULT_MANIFEST))
    subparsers = parser.add_subparsers(dest="command", required=True)
    subparsers.add_parser("validate-manifest")
    subparsers.add_parser("list-runnable")
    plan = subparsers.add_parser("plan-runs")
    plan.add_argument("--output-root", default=str(EXPECTED_OUTPUT_ROOT))
    inspect = subparsers.add_parser("inspect-entry")
    inspect.add_argument("--model-id", required=True)
    inspect.add_argument("--output-root", default=str(EXPECTED_OUTPUT_ROOT))
    readiness = subparsers.add_parser("readiness")
    readiness.add_argument("--output-root", default=str(EXPECTED_OUTPUT_ROOT))
    readiness.add_argument("--report-path", required=True)
    readiness.add_argument("--evidence-path", required=True)
    aggregate_parser = subparsers.add_parser("aggregate")
    aggregate_parser.add_argument(
        "--output-root", default=str(EXPECTED_OUTPUT_ROOT)
    )
    aggregate_parser.add_argument("--report-path", required=True)
    aggregate_parser.add_argument("--evidence-path", required=True)
    aggregate_parser.add_argument("--require-complete", action="store_true")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        manifest = load_manifest(args.manifest)
        if args.command == "validate-manifest":
            print(
                json.dumps(
                    {
                        "status": "PASS",
                        "scope_id": E5_SCOPE27_ID,
                        "counts": manifest["counts"],
                    },
                    ensure_ascii=False,
                )
            )
            return 0
        if args.command == "list-runnable":
            for entry in manifest["entries"]:
                if entry["entry_type"] == "REFERENCE_ONLY_FORMAL_A8":
                    continue
                device = "cpu" if entry["command"] == "evaluate-only" else "cuda"
                print(
                    "\t".join(
                        (
                            entry["model_id"],
                            entry["command"],
                            entry["e5_run_id"],
                            device,
                        )
                    )
                )
            return 0
        if args.command == "plan-runs":
            output_root = Path(args.output_root).resolve()
            if output_root != EXPECTED_OUTPUT_ROOT.resolve():
                raise ScopeGateError("Formal output root mismatch.")
            for entry in manifest["entries"]:
                if entry["entry_type"] == "REFERENCE_ONLY_FORMAL_A8":
                    continue
                result = inspect_run(manifest, entry, output_root)
                if result["ready"]:
                    action = "SKIP_COMPLETED_IDENTITY_MATCH"
                    reason = "COMPLETED_IDENTITY_MATCH"
                elif not result["found"]:
                    action = "RUN"
                    reason = "RUN_NOT_FOUND"
                else:
                    action = "BLOCK_EXISTING_PRESERVED"
                    reason = "|".join(result["reasons"])
                device = "cpu" if entry["command"] == "evaluate-only" else "cuda"
                print(
                    "\t".join(
                        (
                            entry["model_id"],
                            entry["command"],
                            entry["e5_run_id"],
                            device,
                            action,
                            reason,
                        )
                    )
                )
            return 0
        if args.command == "inspect-entry":
            entry = _entry(manifest, args.model_id)
            if entry["entry_type"] == "REFERENCE_ONLY_FORMAL_A8":
                result = inspect_a8(manifest, entry)
            else:
                result = inspect_run(
                    manifest, entry, Path(args.output_root).resolve()
                )
            print(json.dumps(result, ensure_ascii=False, indent=2))
            if result["ready"]:
                return 0
            if not result["found"]:
                return 10
            return 20
        if args.command == "readiness":
            report = build_readiness(
                manifest, Path(args.output_root).resolve()
            )
            write_readiness_outputs(
                report,
                Path(args.report_path).resolve(),
                Path(args.evidence_path).resolve(),
            )
            print(json.dumps(report, ensure_ascii=False, indent=2))
            return 0 if report["status"] == "READY" else 4
        if args.command == "aggregate":
            result = aggregate(
                manifest,
                Path(args.output_root).resolve(),
                require_complete=args.require_complete,
                report_path=Path(args.report_path).resolve(),
                evidence_path=Path(args.evidence_path).resolve(),
            )
            print(json.dumps(result, ensure_ascii=False, indent=2))
            return 0
    except (OSError, ValueError, KeyError, ScopeGateError) as exc:
        print(f"ERROR: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 2
    return 2


if __name__ == "__main__":
    raise SystemExit(main())
