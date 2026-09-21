from pathlib import Path

from docx import Document


INPUT = Path(r"C:\Users\12811\Desktop\手稿.docx")
OUTPUT = Path(r"C:\Users\12811\Desktop\手稿_消融实验A7调整版.docx")


def iter_story_paragraphs(doc):
    """Yield paragraphs in the main story, tables, headers, and footers."""
    for paragraph in doc.paragraphs:
        yield paragraph

    def walk_table(table):
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph
                for nested in cell.tables:
                    yield from walk_table(nested)

    for table in doc.tables:
        yield from walk_table(table)

    for section in doc.sections:
        for part in (section.header, section.first_page_header, section.even_page_header,
                     section.footer, section.first_page_footer, section.even_page_footer):
            for paragraph in part.paragraphs:
                yield paragraph
            for table in part.tables:
                yield from walk_table(table)


def replace_in_runs(doc, old, new):
    count = 0
    for paragraph in iter_story_paragraphs(doc):
        for run in paragraph.runs:
            if old in run.text:
                count += run.text.count(old)
                run.text = run.text.replace(old, new)
    return count


def replace_paragraph_text(paragraph, text):
    """Keep the paragraph style and first-run formatting while replacing its text."""
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def remove_row_with_label(table, label):
    matches = [row for row in table.rows if row.cells and row.cells[0].text.strip() == label]
    if len(matches) != 1:
        raise RuntimeError(f"Expected exactly one row labeled {label!r}, found {len(matches)}")
    tr = matches[0]._tr
    tr.getparent().remove(tr)


def visible_text(doc):
    parts = []
    for paragraph in iter_story_paragraphs(doc):
        parts.append(paragraph.text)
    return "\n".join(parts)


def main():
    if not INPUT.exists():
        raise FileNotFoundError(INPUT)

    doc = Document(str(INPUT))

    # The old A7 is the Temporal-Granularity Prompt row in the E8 table.
    # Remove that result row before renaming the old A8 loss-ablation row.
    if len(doc.tables) < 77:
        raise RuntimeError(f"Unexpected table count: {len(doc.tables)}")
    remove_row_with_label(doc.tables[74], "A7")

    # Rename every visible reference to the former A8.  Explicit paragraph
    # rewrites below also remove the former A7's prose and statistics.
    replaced = replace_in_runs(doc, "A8", "A7")
    if replaced < 1:
        raise RuntimeError("No A8 references were found to rename")
    replace_in_runs(doc, "A4-A7", "A4-A6")
    replace_in_runs(doc, "A4–A7", "A4–A6")

    # Keep the manuscript's existing structure and formatting while making
    # the ablation inventory and its interpretation consistent with A0-A7.
    replacements = {
        150: (
            "A0-A7 均在 fixed_dual、H64、C1 和相同评价流程下比较。"
            "A1-A3 与 A7 保持原有组件消融定义；A4 定义为 Mean-Pooling Macro Prompt，"
            "仅将 Macro Prompt 的时间聚合由 attention pooling 改为等权 mean pooling，"
            "并保持 macro_prompt_len=4、Macro Prompt 开启及后续结构不变；A5 继续采用 "
            "Short-context Reverse Cross；A6 继续采用 Additive Cross Fusion。"
            "A4 的 effective_config_diff 显示唯一差异字段为 macro_prompt_pooling，"
            "且该 run 完成正式评估并通过协议审计。A3 的 w/o Diffusion 同时采用 Simple graph；"
            "A7 将损失由 msmg_dwu_loss 切换为 masked_score_aligned_hybrid，因此 A7 表示 "
            "MS-MG-DWU 联合目标的关闭/替代，而不是只屏蔽一个网络层。"
        ),
        154: (
            "A0 的平均 Score 为 9.7247（表内四舍五入为 9.725）。按平均 Score 排序，"
            "A0 为 9.7247、A6 为 9.9790、A2 为 10.2220、A1 为 10.2317、"
            "A4 为 10.3267、A5 为 10.3347、A7 为 10.3637、A3 为 10.7410。"
            "A4 相对 A0 增加 0.6020（6.19%），A7 增加 0.6390（6.57%）；"
            "因此在当前单 seed 结果中所有消融变体平均 Score 均高于 A0。"
        ),
        158: (
            "A5 继续采用 Short-context Reverse Cross，H3/H6/H10 Score 相对 A0 分别变化 "
            "0.849/0.568/0.413，平均增加 0.6100（6.27%），三个 horizon 均退化。"
        ),
        163: (
            "E7 将 A0 与三个空间机制消融进行内部对照：A1 去除 Spatial Graph，A2 去除 Adaptive Graph，"
            "A3 去除 Diffusion；同时与 GCN、STGCN、DCRNN、MTGNN、STID 和 TSMixer 进行端到端结果比较。"
            "原始 E7 工作簿保留这组 A0–A3 与外部模型的核心表，不将 A4–A6 混入 E7 的图机制核心排名。"
            "表中统一列出 H3、H6 和 H10 下的 Score、MAE、RMSE 与 R²。"
        ),
        166: (
            "A4-A6 的 Prompt/Cross-Fusion 正式结果归入 E8 表 4-4 的上半部分，不纳入 E7 的图机制核心表。"
            "E7 仍按原始 E7 工作簿的 A0-A3 与外部图机制对照口径报告；E8 上半部分的 A0、A4-A6 "
            "按最新正式定义统一更新。A4、A5、A6 的平均 Score 分别为 10.3267、10.3347、9.9790，"
            "相对 A0 的平均变化分别为 +0.6020、+0.6100 和 +0.2543；A4 通过 pooling strategy 的"
            "单字段 effective_config_diff 审计。"
        ),
        168: (
            "E8 聚焦跨粒度交互与 Prompt 机制。根据更新版配置审计，A4-A6 均为单因素受控比较："
            "A4 为 Mean-Pooling Macro Prompt，仅将 Macro Prompt 的 attention temporal pooling 改为 mean pooling，"
            "保持 macro_prompt_len=4、Prompt 容量、后续融合和解码路径不变；A5 继续采用 Short-context Reverse Cross；"
            "A6 继续采用 Additive Cross Fusion。同时与 PatchTST、iTransformer、TimeXer、MultiPatchFormer、"
            "TimeMixer 和 TimeFilter 进行外部端到端比较。"
        ),
        170: (
            "表 4-4 的内部结果显示，A4 的平均 Score 为 10.3267，相对 A0 增加 0.6020（6.19%）；"
            "A5、A6 的平均增加分别为 0.6100 和 0.2543。A4 的差异严格对应 attention 与 mean 时间聚合；"
            "A4-A6 均不再解释为关闭整个 Prompt 模块或改变 Prompt 容量。完整模型的作用按多视界整体权衡解释，"
            "不能把任一单次消融写成跨 seed 的统计结论。"
        ),
        198: (
            "A4-A6 的当前正式配置在多视界上表现不同。A4 将 Macro Prompt 的 attention temporal pooling 改为 "
            "mean pooling，保持 macro_prompt_len=4 及其余结构不变，平均 Score 增加 0.6020（6.19%），"
            "H3/H6/H10 分别增加 0.902/0.590/0.314。A5 继续采用 Short-context Reverse Cross，平均 Score 增加 "
            "0.6100（6.27%），三个 horizon 均退化。A6 继续采用 Additive Cross Fusion，平均 Score 增加 "
            "0.2543（2.62%），三个 horizon 均变差。"
        ),
        199: (
            "二级指标也不能脱离官方 Score 单独下结论。A4 的 H10 RMSE/R² 为 165.844/0.842，相比 A0 的 "
            "159.341/0.854 同向退化；A5 和 A6 的 H10 RMSE/R² 分别为 164.419/0.845 和 158.890/0.855，"
            "其中 A5 同向退化，A6 出现局部改善，但官方 H10 Score 仍分别增加 0.413 和 0.081。本文因此固定采用"
            "官方 Score 进行主排序，同时报告 MAE、RMSE 和 R²，保留指标之间的局部不一致。"
        ),
        211: (
            "另外，A4 的 attention-versus-mean temporal pooling、A7 的损失替代以及 P4/P5 的 runner 状态冲突"
            "均在结果中显式保留。E7 核心表与 E8 表上半部分的 A4-A6 结果属于不同的内部实验口径，E9 的六模型"
            "迁移结果属于成对迁移口径，外部比较仍属于跨结果源的端到端比较；A4 的新结果是单一 pooling 机制的"
            "结构受控对照，A7 的新结果是损失替代的整体效应，不将这些差异扩展解释为内部组件的纯因果贡献。"
            "E9 每对 run 的初始状态字节级一致性尚未由源审计验证；volatility、Ramp、Shared difficult Top10%、"
            "E5 与多随机种子统计检验仍待补充，相关性、可解释性和因果作用之间保持谨慎区分。"
        ),
    }
    for index, text in replacements.items():
        replace_paragraph_text(doc.paragraphs[index], text)

    # The full ablation table now contains A0-A7; its former A8 row was
    # renamed by the global pass above.  The E8 table contains A0 and A4-A6.
    expected_full = [doc.tables[72].rows[i].cells[0].text.strip() for i in range(2, len(doc.tables[72].rows))]
    if expected_full != ["A0", "A1", "A2", "A3", "A4", "A5", "A6", "A7"]:
        raise RuntimeError(f"Unexpected full ablation labels: {expected_full}")
    expected_e8 = [doc.tables[74].rows[i].cells[0].text.strip() for i in range(2, 6)]
    if expected_e8 != ["A0", "A4", "A5", "A6"]:
        raise RuntimeError(f"Unexpected E8 internal labels: {expected_e8}")

    text = visible_text(doc)
    forbidden = [
        "A0-A8", "A0–A8", "A1–A8", "A4-A7", "A4–A7", "Temporal-Granularity",
        "node identity", "st_prompt_use_node_identity", "9.9690", "0.2443",
        "0.534/0.170/0.029", "158.129/0.857", "8.318",
    ]
    leftovers = [item for item in forbidden if item in text]
    if leftovers:
        raise RuntimeError(f"Former A7 content or old ranges remain: {leftovers}")
    if "A7 将损失由 msmg_dwu_loss" not in text:
        raise RuntimeError("The renamed A7 loss-ablation description is missing")

    doc.save(str(OUTPUT))
    print(f"Saved {OUTPUT}")
    print(f"Renamed visible A8 references: {replaced}")
    print(f"Full ablation labels: {expected_full}")
    print(f"E8 internal labels: {expected_e8}")


if __name__ == "__main__":
    main()
