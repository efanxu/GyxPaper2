from __future__ import annotations

import csv
import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[3]
SOURCE_ROOT = PROJECT_ROOT / "custom_models/src"
for selected in (PROJECT_ROOT, SOURCE_ROOT):
    if str(selected) not in sys.path:
        sys.path.insert(0, str(selected))

import torch

from benchmark_v2.experiments.e5_common_loss.a8_reference import create_a8_reference
from benchmark_v2.experiments.e5_common_loss.aggregation import aggregate
from benchmark_v2.experiments.e5_common_loss.readiness import CORE_EXPECTED, build_readiness
from st_mgprompt.a8_batch4_contract import A8_REFERENCE_RELATIVE_PATH
from st_mgprompt.a8_batch4_readiness import A8_EXPLICIT_CONFIG
from scripts import e5_batch4_scope27_gate as gate


def _write_json(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, allow_nan=True), encoding="utf-8")


def _metrics(root: Path, *, nonfinite: bool = False) -> None:
    for horizon in (3, 6, 10):
        _write_json(root / f"metrics_eval_h{horizon}.json", {
            "horizon": horizon,
            "MAE": float("nan") if nonfinite and horizon == 3 else 1.0,
            "RMSE": 1.5,
            "R2": 0.25,
            "Score": 2.0,
            "valid_target_count": 10,
        })


def _assert_no_identity_gates(test: unittest.TestCase, value: object) -> None:
    forbidden = ("ha" + "sh", "sha" + "256", "dig" + "est", "finger" + "print")
    if isinstance(value, dict):
        for key, child in value.items():
            test.assertFalse(any(term in str(key).casefold() for term in forbidden), key)
            _assert_no_identity_gates(test, child)
    elif isinstance(value, list):
        for child in value:
            _assert_no_identity_gates(test, child)


class E5Batch4EndToEndTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.output_root = self.root / "results" / "common_loss_architecture_seed2026"
        self.a8_root = self.root / "results" / "a8" / "component_ablation_a8_bs4_seed2026" / "STMGPrompt_ComponentAblation"
        self.reference_path = self.root / "logs" / "uniform_bs4" / "audit" / "e5_scope27" / "E5_A8_BATCH4_REFERENCE.json"
        self.manifest = gate.load_manifest()
        self._create_a8()
        self._create_e5()

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def _create_a8(self) -> None:
        self.a8_root.mkdir(parents=True, exist_ok=True)
        _write_json(self.a8_root / "effective_config.json", A8_EXPLICIT_CONFIG)
        _write_json(self.a8_root / "run_status.json", {"status": "COMPLETED", "exit_code": 0})
        _metrics(self.a8_root)
        torch.save({"model_state_dict": {}}, self.a8_root / "best_checkpoint.pt")
        create_a8_reference(self.reference_path, run_root=self.a8_root)

    def _create_e5(self) -> None:
        self.output_root.mkdir(parents=True, exist_ok=True)
        for entry in self.manifest["entries"]:
            if entry["entry_type"] == "REFERENCE_ONLY_FORMAL_A8":
                continue
            run_dir = self.output_root / entry["e5_run_id"]
            run_dir.mkdir(parents=True, exist_ok=True)
            effective = {
                **CORE_EXPECTED,
                "model_id": entry["model_id"],
                "run_id": entry["e5_run_id"],
                "loss_id": entry["loss_id"],
                "formal_training": entry["entry_type"] == "TRAIN_COMMON_LOSS",
                "parameter_count": 1,
                "trainable_parameter_count": 1,
            }
            _write_json(run_dir / "effective_config.json", effective)
            _write_json(run_dir / "run_status.json", {"status": "COMPLETED", "exit_code": 0})
            _metrics(run_dir)
            if entry["entry_type"] == "TRAIN_COMMON_LOSS":
                torch.save({"model_state_dict": {}}, run_dir / "best_checkpoint.pt")

    def _readiness(self) -> dict[str, object]:
        return build_readiness(
            manifest=self.manifest,
            output_root=self.output_root,
            a8_reference_path=self.reference_path,
            a8_run_root=self.a8_root,
        )

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl is required for XLSX E2E verification")
    def test_full_24_plus_2_plus_1_readiness_and_aggregate(self) -> None:
        from openpyxl import load_workbook
        tracked_before = subprocess.run(
            ["git", "diff", "--name-only"], check=True, capture_output=True, text=True,
        ).stdout
        common = self._readiness()
        gated = gate.build_readiness(
            self.manifest, self.output_root,
            a8_reference_path=self.reference_path, a8_run_root=self.a8_root,
        )
        self.assertEqual(common["status"], "READY")
        self.assertEqual(gated["status"], "READY")
        self.assertEqual(common["ready_entries"], 27)
        result = aggregate(
            output_root=self.output_root,
            require_complete=True,
            manifest=self.manifest,
            a8_reference_path=self.reference_path,
            a8_run_root=self.a8_root,
        )
        self.assertEqual(result["row_count"], 27)
        csv_path = self.output_root / "common_loss_architecture_seed2026.csv"
        with csv_path.open("r", encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual(len(rows), 27)
        self.assertEqual(sum(row["model_id"] == "st_mgprompt_a8" for row in rows), 1)
        self.assertTrue(all(
            row["run_id"].endswith("_bs4_seed2026")
            for row in rows if row["model_id"] != "st_mgprompt_a8"
        ))
        benchmark_ids = [entry["e5_run_id"] for entry in self.manifest["entries"][:-1]]
        self.assertTrue(all(run_id.endswith("_bs4_seed2026") for run_id in benchmark_ids))
        workbook = load_workbook(self.output_root / "COMMON_LOSS_ARCHITECTURE_SEED2026.xlsx", read_only=True)
        self.assertEqual(workbook.sheetnames, ["Trainable structures", "Non-trainable references", "All 27"])
        self.assertEqual(workbook["Trainable structures"].max_row, 25)
        self.assertEqual(workbook["Non-trainable references"].max_row, 4)
        self.assertEqual(workbook["All 27"].max_row, 28)
        workbook.close()
        audit = json.loads((self.output_root / "common_loss_architecture_audit.json").read_text(encoding="utf-8"))
        reference = json.loads(self.reference_path.read_text(encoding="utf-8"))
        _assert_no_identity_gates(self, audit)
        _assert_no_identity_gates(self, reference)
        tracked_after = subprocess.run(
            ["git", "diff", "--name-only"], check=True, capture_output=True, text=True,
        ).stdout
        self.assertEqual(tracked_after, tracked_before)
        self.assertTrue(A8_REFERENCE_RELATIVE_PATH.startswith("custom_models/logs/"))

    def test_invalid_explicit_metadata_artifacts_block_readiness_and_aggregate(self) -> None:
        first = self.manifest["entries"][0]
        run_dir = self.output_root / first["e5_run_id"]
        config_path = run_dir / "effective_config.json"
        metric_path = run_dir / "metrics_eval_h3.json"
        original_config = json.loads(config_path.read_text(encoding="utf-8"))
        original_metric = json.loads(metric_path.read_text(encoding="utf-8"))
        cases = (
            ("missing_scope_id", {key: value for key, value in original_config.items() if key != "scope_id"}, original_metric, "ARCHIVE_AND_RUN"),
            ("missing_run_id", {key: value for key, value in original_config.items() if key != "run_id"}, original_metric, "ARCHIVE_AND_RUN"),
            ("wrong_loss_id", {**original_config, "loss_id": "masked_mse"}, original_metric, "BLOCK_EXPLICIT_CONFIG_CONFLICT"),
            ("batch32_profile", {**original_config, "training_batch_profile_id": "default_batch32"}, original_metric, "BLOCK_EXPLICIT_CONFIG_CONFLICT"),
            ("nonfinite_metric", original_config, {**original_metric, "MAE": float("nan")}, "ARCHIVE_AND_RUN"),
        )
        for name, config, metric, expected_action in cases:
            with self.subTest(name=name):
                _write_json(config_path, config)
                _write_json(metric_path, metric)
                common = self._readiness()
                gated = gate.build_readiness(
                    self.manifest, self.output_root,
                    a8_reference_path=self.reference_path, a8_run_root=self.a8_root,
                )
                self.assertEqual(common["status"], "NOT_READY")
                self.assertEqual(gated["status"], common["status"])
                self.assertEqual(gated["ready_entries"], common["ready_entries"])
                inspected = gate.inspect_run(self.manifest, first, self.output_root)
                self.assertEqual(gate._plan_action(first, inspected), expected_action)
                with self.assertRaises(RuntimeError):
                    aggregate(
                        output_root=self.output_root, require_complete=True,
                        manifest=self.manifest,
                        a8_reference_path=self.reference_path, a8_run_root=self.a8_root,
                    )
                _write_json(config_path, original_config)
                _write_json(metric_path, original_metric)

        self.assertEqual(self._readiness()["status"], "READY")
        (self.a8_root / "best_checkpoint.pt").write_bytes(b"not a checkpoint")
        self.assertEqual(self._readiness()["status"], "NOT_READY")


if __name__ == "__main__":
    unittest.main()
