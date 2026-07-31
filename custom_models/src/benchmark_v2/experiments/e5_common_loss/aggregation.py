from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from ...artifacts import atomic_write_json
from ...runtime import PROJECT_ROOT
from ...training_profiles import batch_identity
from .contracts import FORMAL_OUTPUT_ROOT_RELATIVE
from .readiness import build_readiness
from .variant_manifest import build_variant_manifest


def _load(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def aggregate(
    *,
    output_root: str | Path | None = None,
    require_complete: bool = True,
    training_profile: str | None = None,
) -> dict[str, Any]:
    if not require_complete:
        raise ValueError("E5 final aggregation requires --require-complete")
    root = Path(output_root or (PROJECT_ROOT / FORMAL_OUTPUT_ROOT_RELATIVE))
    readiness = build_readiness(
        output_root=root, training_profile=training_profile
    )
    if readiness["status"] != "READY":
        raise RuntimeError(
            f"E5_RESULT_NOT_READY: {readiness['ready_entries']}/29 entries ready"
        )
    manifest = build_variant_manifest(training_profile=training_profile)
    rows = []
    for entry in manifest["entries"]:
        if entry["entry_type"] == "REFERENCE_ONLY_FORMAL_A8":
            from .a8_reference import validate_a8_reference

            reference = validate_a8_reference(
                training_profile=training_profile
            )
            metrics_paths = [Path(path) for path in reference["source_metrics_paths"]]
            effective = _load(Path(reference["source_config_path"]))
            run_status = {"status": "REFERENCE_ONLY"}
        else:
            run_dir = root / entry["e5_run_id"]
            metrics_paths = [run_dir / f"metrics_eval_h{h}.json" for h in (3, 6, 10)]
            effective = _load(run_dir / "effective_config.json")
            run_status = _load(run_dir / "run_status.json")
        metrics = {int(_load(path)["horizon"]): _load(path) for path in metrics_paths}
        row = {
            "model": entry["display_name"],
            "model_id": entry["model_id"],
            "category": entry.get("category"),
            "entry_type": entry["entry_type"],
            "training_mode": entry["training_mode"],
            "loss_id": entry["loss_id"],
            "loss_profile_hash": entry["loss_profile_hash"],
            "training_loss": entry.get("training_loss", entry["loss_id"]),
            "common_loss_evaluation_applied": entry.get(
                "common_loss_evaluation_applied", False
            ),
            "best_epoch": effective.get("best_epoch"),
            "parameter_count": effective.get("parameter_count"),
            "trainable_parameter_count": effective.get("trainable_parameter_count"),
            "training_status": run_status.get("status"),
            "preflight_status": effective.get("preflight_identity", {}).get("status")
            if isinstance(effective.get("preflight_identity"), dict)
            else None,
            **batch_identity(training_profile),
        }
        for horizon in (3, 6, 10):
            for metric in ("Score", "MAE", "RMSE", "R2"):
                row[f"H{horizon}_{metric}"] = metrics[horizon][metric]
        row["average_Score"] = sum(
            float(metrics[horizon]["Score"]) for horizon in (3, 6, 10)
        ) / 3.0
        rows.append(row)
    fieldnames = list(rows[0])
    csv_path = root / "common_loss_architecture_seed2026.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    md_path = root / "COMMON_LOSS_ARCHITECTURE_SEED2026.md"
    md_lines = [
        "# Common-loss architecture seed 2026",
        "",
        "Current results use one seed (2026); multi-seed significance is deferred to E11.",
        "",
        "| model | mode | H3 Score | H6 Score | H10 Score | average Score |",
        "|---|---:|---:|---:|---:|---:|",
    ]
    for row in rows:
        md_lines.append(
            f"| {row['model']} | {row['training_mode']} | "
            f"{row['H3_Score']:.6f} | {row['H6_Score']:.6f} | "
            f"{row['H10_Score']:.6f} | {row['average_Score']:.6f} |"
        )
    md_path.write_text("\n".join(md_lines) + "\n", encoding="utf-8")
    audit_path = root / "common_loss_architecture_audit.json"
    atomic_write_json(
        audit_path,
        {
            "status": "PASS",
            "require_complete": True,
            "readiness": readiness,
            "row_count": len(rows),
            **batch_identity(training_profile),
        },
    )
    xlsx_path = root / "COMMON_LOSS_ARCHITECTURE_SEED2026.xlsx"
    try:
        from openpyxl import Workbook

        workbook = Workbook()
        trained = workbook.active
        trained.title = "Trainable structures"
        references = workbook.create_sheet("Non-trainable references")
        appendix = workbook.create_sheet("All 29")
        for sheet, selected in (
            (
                trained,
                [
                    row
                    for row in rows
                    if row["entry_type"] != "EVALUATE_ONLY_COMMON_LOSS_DIAGNOSTIC"
                ],
            ),
            (
                references,
                [
                    row
                    for row in rows
                    if row["entry_type"] == "EVALUATE_ONLY_COMMON_LOSS_DIAGNOSTIC"
                ],
            ),
            (appendix, rows),
        ):
            sheet.append(fieldnames)
            for row in selected:
                sheet.append([row.get(field) for field in fieldnames])
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
        workbook.save(xlsx_path)
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for complete E5 aggregation") from exc
    return {
        "status": "PASS",
        "row_count": len(rows),
        "files": [str(xlsx_path), str(md_path), str(csv_path), str(audit_path)],
        **batch_identity(training_profile),
    }
